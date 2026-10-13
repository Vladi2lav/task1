"""Парсер биржи ETS (ets.kz): страницы «итоги торгов → год → месяц» и XLSX-файлы за каждый торговый день."""
from __future__ import annotations

import io
import logging
import re
from dataclasses import dataclass
from datetime import date
from typing import Iterator, List, Optional
from urllib.parse import urljoin, urlparse

import openpyxl
from bs4 import BeautifulSoup
from django.conf import settings

from .base import BaseScraper, ScraperError, TradeRecord
from .fingerprint import assign_occurrences
from .http import HttpClient
from .normalize import (
    clean_text, normalize_bin, normalize_company_name, parse_date, parse_decimal, split_bilingual_product,
)

logger = logging.getLogger(__name__)

# основы названий месяцев -> номер месяца (слова «октябрь», «Октября» и т.п.)
_MONTH_STEMS = {
    'январ': 1, 'феврал': 2, 'март': 3, 'апрел': 4, 'мая': 5, 'май': 5, 'июн': 6,
    'июл': 7, 'август': 8, 'сентябр': 9, 'октябр': 10, 'ноябр': 11, 'декабр': 12,
}

# колонки листа определяются по тексту шапки: поле -> регулярные выражения, которые должны совпасть все
# (порядок важен: более конкретные правила идут раньше)
_COLUMN_RULES = [
    ('buyer_broker', (r'брокер', r'покупател')),
    ('seller_broker', (r'брокер', r'продав')),
    ('buyer_bin', (r'\bбин\b', r'покупател')),
    ('seller_bin', (r'\bбин\b', r'продав')),
    ('buyer_name', (r'наименование', r'покупател')),
    ('seller_name', (r'наименование', r'продав')),
    ('product_name', (r'наименование', r'товар')),
    ('tnved', (r'вэд',)),
    ('product_code', (r'код', r'товар')),
    ('deals_count', (r'кол(ичество|-во)', r'сделок')),
    ('price_open', (r'цена', r'открыт')),
    ('price_close', (r'цена', r'закрыт')),
    ('price_max', (r'макс',)),
    ('price_min', (r'\bмин',)),
    ('price_avg', (r'средневзв',)),
    ('amount', (r'общая', r'сумм')),
]
# сумма сделки обязательной не считается: в файлах 2025 года её нет
_REQUIRED_COLUMNS = {'buyer_name', 'seller_name', 'product_name'}
_HEADER_SEARCH_ROWS = 30
_DATE_IN_TEXT = re.compile(r'\d{1,2}\.\d{1,2}\.\d{4}')
# базис поставки (Инкотермс) в конце названия товара: «..., EXW»
_BASIS_AT_END = re.compile(r'(?:^|[\s,])(EXW|FCA|FAS|FOB|CFR|CIF|CPT|CIP|DAP|DPU|DDP|DAT|DAF)\s*$')


@dataclass
class _XlsxLink:
    url: str
    trade_date: Optional[date]  # по названию файла; None, если в названии даты нет


def parse_period_label(text: str):
    """«Итоги торгов 2026 год Октябрь» -> (2026, 10); без месяца -> (2025, None); без года -> (None, None)."""
    year_match = re.search(r'\b(20\d{2})\b', text)
    year = int(year_match.group(1)) if year_match else None
    month = None
    for word in re.findall(r'[а-яё]+', text.lower()):
        month = next((number for stem, number in _MONTH_STEMS.items() if word.startswith(stem)), None)
        if month:
            break
    return year, month


class EtsScraper(BaseScraper):
    exchange_code = 'ets'
    exchange_name = 'ETS'

    def __init__(self, http: Optional[HttpClient] = None, config: Optional[dict] = None):
        config = config or settings.SCRAPER_CONFIG['sources']['ets']
        self.base_url = config['base_url']
        self.results_url = urljoin(self.base_url, config['results_path'])
        self.website = self.base_url
        self.http = http or HttpClient()
        self.errors: List[str] = []

    # --- основной проход ---------------------------------------------------

    def fetch(self, date_from: date, date_to: date) -> Iterator[TradeRecord]:
        self.errors = []
        for month_url in self._month_pages(date_from, date_to):
            try:
                links = self.parse_xlsx_links(self.http.get_text(month_url))
            except ScraperError as exc:
                self._report_error(f'Страница месяца {month_url}: {exc}')
                continue
            for link in links:
                if link.trade_date and not date_from <= link.trade_date <= date_to:
                    continue
                records = self._load_file(link, date_from, date_to)
                if records:
                    logger.info('ETS %s: %d строк', link.trade_date or link.url, len(records))
                yield from records

    def _report_error(self, message: str) -> None:
        logger.error('ETS: %s', message)
        self.errors.append(message)

    def _load_file(self, link: _XlsxLink, date_from: date, date_to: date) -> List[TradeRecord]:
        """Скачать и разобрать один файл. Ошибка — в self.errors, файл пропускается целиком."""
        try:
            content = self.http.get_bytes(link.url)
            records = self.parse_xlsx(content, link.url, link.trade_date)
        except Exception as exc:  # один плохой файл не должен ронять весь сбор
            self._report_error(f'Файл {link.url}: {exc}')
            return []
        return [r for r in records if date_from <= r.trade_date <= date_to]

    # --- обход страниц сайта ----------------------------------------------

    def _month_pages(self, date_from: date, date_to: date) -> Iterator[str]:
        """Ссылки на страницы месяцев, попадающих в период (скачивать лишнее не нужно)."""
        first, last = (date_from.year, date_from.month), (date_to.year, date_to.month)
        try:
            index_html = self.http.get_text(self.results_url)
        except ScraperError as exc:
            raise ScraperError(f'Не открылась страница итогов торгов ETS: {exc}') from exc

        year_urls = self.parse_year_links(index_html)
        if not year_urls:
            raise ScraperError('На странице итогов торгов ETS не найдены ссылки на годы: изменилась вёрстка?')

        for year, year_url in sorted(year_urls.items()):
            if not first[0] <= year <= last[0]:
                continue
            try:
                months = self.parse_month_links(self.http.get_text(year_url))
            except ScraperError as exc:
                self._report_error(f'Страница года {year}: {exc}')
                continue
            for (link_year, month), month_url in sorted(months.items()):
                if first <= (link_year, month) <= last:
                    yield month_url

    def _page_links(self, html: str):
        """Ссылки внутри раздела итогов торгов: (текст, абсолютный URL, путь)."""
        results_path = urlparse(self.results_url).path
        for a in BeautifulSoup(html, 'lxml').find_all('a', href=True):
            url = urljoin(self.results_url, a['href'])
            path = urlparse(url).path
            if path.startswith(results_path) and path != results_path:
                yield clean_text(a.get_text(' ')), url

    def parse_year_links(self, html: str) -> dict:
        """Главная страница итогов: {год: URL страницы года}."""
        years = {}
        for text, url in self._page_links(html):
            year, month = parse_period_label(text)
            if year and not month:
                years.setdefault(year, url)
        return years

    def parse_month_links(self, html: str) -> dict:
        """Страница года: {(год, месяц): URL страницы месяца}. Месяц определяется по тексту ссылки."""
        months = {}
        for text, url in self._page_links(html):
            year, month = parse_period_label(text)
            if year and month:
                months.setdefault((year, month), url)
        return months

    def parse_xlsx_links(self, html: str) -> List[_XlsxLink]:
        """Страница месяца: ссылки на XLSX действующих редакций файлов."""
        soup = BeautifulSoup(html, 'lxml')
        links, seen = [], set()
        for a in soup.select('a.i_documents_icon[href]'):
            url = urljoin(self.base_url, a['href'])
            if not url.lower().endswith('.xlsx') or url in seen:
                continue
            if not self._is_current_edition(a):
                logger.debug('ETS: пропущена не действующая редакция %s', url)
                continue
            seen.add(url)
            links.append(_XlsxLink(url, self._date_of_link(a)))
        return links

    @staticmethod
    def _is_current_edition(link) -> bool:
        """Файл лежит в блоке «Действующая редакция» (архивные редакции пропускаются)."""
        block = link.find_parent(class_='i_documents_section_items')
        title = block.find(class_='i_documents_section_list_title') if block else None
        return title is None or 'действующ' in title.get_text().lower()

    @staticmethod
    def _date_of_link(link) -> Optional[date]:
        """Дата торгов из названия файла (download), заголовка документа или раздела."""
        candidates = [link.get('download', '')]
        item = link.find_parent(class_='i_documents_item')
        if item and item.find(class_='i_documents_item_title'):
            candidates.append(item.find(class_='i_documents_item_title').get_text())
        section = link.find_parent(class_='i_documents_section')
        if section and section.find(class_='i_documents_section_name'):
            candidates.append(section.find(class_='i_documents_section_name').get_text())
        for text in candidates:
            match = _DATE_IN_TEXT.search(text)
            if match:
                return parse_date(match.group())
        return None

    # --- разбор XLSX -------------------------------------------------------

    def parse_xlsx(self, content: bytes, source_url: str, file_date: Optional[date] = None) -> List[TradeRecord]:
        """Строки всех листов файла в виде TradeRecord. Строка «Итого» пропускается.

        В некоторых файлах название и код ТН ВЭД товара — формулы VLOOKUP по скрытому листу-справочнику.
        Формулы читаются вместе с их сохранёнными значениями (data_only=True); если значения нет,
        товар берётся из справочника по коду.
        """
        stream = io.BytesIO(content)
        workbook = openpyxl.load_workbook(stream, data_only=False)
        cached = openpyxl.load_workbook(stream, data_only=True)
        directory = self._product_directory(cached)

        records: List[TradeRecord] = []
        for sheet in workbook.worksheets:
            records += self._parse_sheet(sheet, cached[sheet.title], directory, source_url, file_date)
        if not records:
            raise ScraperError('в файле не найдено ни одной строки торгов')
        return assign_occurrences(records)

    def _product_directory(self, workbook) -> dict:
        """Справочник товаров из листа без покупателей/продавцов: {код товара: (название, ТН ВЭД)}."""
        directory = {}
        for sheet in workbook.worksheets:
            rows = list(sheet.iter_rows(values_only=True))
            header_index, columns = self._find_header(rows, min_columns=3)
            if columns is None or 'buyer_name' in columns or 'product_code' not in columns:
                continue
            for row in rows[header_index + 1:]:
                code = clean_text(row[columns['product_code']])
                if code:
                    tnved = row[columns['tnved']] if 'tnved' in columns else ''
                    directory[code] = (row[columns['product_name']], tnved)
        return directory

    def _parse_sheet(self, sheet, cached_sheet, directory: dict, source_url: str,
                     file_date: Optional[date]) -> List[TradeRecord]:
        # формулы заменяются сохранёнными значениями (None, если Excel их не сохранил)
        rows = [
            tuple(cached if isinstance(value, str) and value.startswith('=') else value
                  for value, cached in zip(row, cached_row))
            for row, cached_row in zip(sheet.iter_rows(values_only=True), cached_sheet.iter_rows(values_only=True))
        ]
        header_index, columns = self._find_header(rows)
        if columns is None:
            logger.debug('ETS: на листе «%s» нет таблицы торгов, лист пропущен', sheet.title)
            return []
        missing = _REQUIRED_COLUMNS - columns.keys()
        if missing:
            raise ScraperError(f'на листе «{sheet.title}» нет колонок: {", ".join(sorted(missing))}')

        title_date = self._title_date(rows[:header_index])
        if file_date and title_date and file_date != title_date:
            logger.warning('ETS: дата в названии файла (%s) и в заголовке листа (%s) не совпадают, берётся дата файла',
                           file_date, title_date)
        trade_date = file_date or title_date
        if trade_date is None:
            raise ScraperError(f'не удалось определить дату торгов на листе «{sheet.title}»')

        records = []
        for row in rows[header_index + 1:]:
            record = self._parse_row(row, columns, directory, trade_date, source_url)
            if record:
                records.append(record)
        return records

    @staticmethod
    def _find_header(rows, min_columns: int = 5) -> tuple:
        """Найти строку шапки и сопоставить колонки полям по тексту шапки. -> (индекс строки, {поле: колонка})."""
        for index, row in enumerate(rows[:_HEADER_SEARCH_ROWS]):
            columns = {}
            for col, cell in enumerate(row):
                text = clean_text(cell).lower()
                if not text:
                    continue
                for field, patterns in _COLUMN_RULES:
                    if field not in columns and all(re.search(p, text) for p in patterns):
                        columns[field] = col
                        break
            if len(columns) >= min_columns and 'product_name' in columns:
                return index, columns
        return -1, None

    @staticmethod
    def _title_date(rows) -> Optional[date]:
        """Дата торгов из заголовка листа («ИТОГИ ТОРГОВ 07.10.2026»)."""
        for row in rows:
            for cell in row:
                match = _DATE_IN_TEXT.search(clean_text(cell))
                if match:
                    return parse_date(match.group())
        return None

    def _parse_row(self, row, columns: dict, directory: dict, trade_date: date,
                   source_url: str) -> Optional[TradeRecord]:
        def cell(field):
            col = columns.get(field)
            return row[col] if col is not None and col < len(row) else None

        if any(clean_text(value).lower().startswith('итого') for value in row if isinstance(value, str)):
            return None  # итоговая строка (в ней формула =SUM)
        product_code = clean_text(cell('product_code'), 50)
        directory_name, directory_tnved = directory.get(product_code, ('', ''))
        product_name = clean_text(split_bilingual_product(cell('product_name') or directory_name), 500)
        if not product_name:
            return None  # пустая или служебная строка
        tnved = clean_text(cell('tnved') or directory_tnved, 30)

        deals = parse_decimal(cell('deals_count'))
        basis = _BASIS_AT_END.search(product_name)

        return TradeRecord(
            exchange_code=self.exchange_code,
            trade_date=trade_date,
            product_name=product_name,
            product_code=product_code,
            tnved=tnved,
            buyer_name=normalize_company_name(cell('buyer_name')),
            buyer_bin=normalize_bin(cell('buyer_bin')),
            seller_name=normalize_company_name(cell('seller_name')),
            seller_bin=normalize_bin(cell('seller_bin')),
            buyer_broker=normalize_company_name(cell('buyer_broker')),
            seller_broker=normalize_company_name(cell('seller_broker')),
            deals_count=int(deals) if deals is not None else None,
            quantity=None,  # ETS объём не публикует
            price_open=parse_decimal(cell('price_open')),
            price_close=parse_decimal(cell('price_close')),
            price_min=parse_decimal(cell('price_min')),
            price_max=parse_decimal(cell('price_max')),
            price_avg=parse_decimal(cell('price_avg')),
            amount=parse_decimal(cell('amount')),
            delivery_basis=basis.group(1) if basis else '',
            source_url=source_url,
        )
