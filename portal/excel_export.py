"""Выгрузка текущей выборки итогов торгов в Excel (xlsx).

Заголовки колонок, названия листов и лист «Параметры» — на языке интерфейса
(язык запроса: cookie django_language / Accept-Language). Данные бирж не переводятся.
"""
from copy import copy
from io import BytesIO

from django.conf import settings
from django.utils import formats, timezone
from django.utils.translation import gettext as _, gettext_lazy

from . import i18n
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# Максимум строк в файле по умолчанию (защита от выгрузки всей базы).
# Можно переопределить в settings.py: PORTAL_EXPORT_MAX_ROWS = 20000.
# Ориентир по времени: openpyxl записывает около 1000 строк в секунду.
EXPORT_MAX_ROWS = 50_000


def export_max_rows() -> int:
    return int(getattr(settings, 'PORTAL_EXPORT_MAX_ROWS', EXPORT_MAX_ROWS))


DATE_FORMAT = 'DD.MM.YYYY'
# порядок дня и месяца в ячейках дат — как принято для языка (en — британский дд/мм/гггг)
DATE_FORMATS = {'en': 'DD/MM/YYYY'}
TIME_FORMAT = 'HH:MM:SS'
PRICE_FORMAT = '#,##0.00'
QUANTITY_FORMAT = '#,##0.000'
INT_FORMAT = '#,##0'

# (заголовок, ширина, функция получения значения из записи Trade, формат ячейки)
COLUMNS = [
    (gettext_lazy('Биржа'), 10, lambda t: i18n.exchange_label(t.exchange.code, t.exchange.name), None),
    (gettext_lazy('Дата торгов'), 13, lambda t: t.trade_date, DATE_FORMAT),
    (gettext_lazy('Время'), 10, lambda t: t.trade_time, TIME_FORMAT),
    (gettext_lazy('Товар'), 50, lambda t: t.product.name, None),
    (gettext_lazy('Код товара'), 14, lambda t: t.product.code, None),
    (gettext_lazy('Код ТН ВЭД'), 14, lambda t: t.product.tnved, None),
    (gettext_lazy('Продавец'), 40, lambda t: t.seller.name if t.seller else '', None),
    (gettext_lazy('БИН продавца'), 15, lambda t: t.seller.bin if t.seller else '', None),
    (gettext_lazy('Покупатель'), 40, lambda t: t.buyer.name if t.buyer else '', None),
    (gettext_lazy('БИН покупателя'), 15, lambda t: t.buyer.bin if t.buyer else '', None),
    (gettext_lazy('Брокер продавца'), 32, lambda t: t.seller_broker.name if t.seller_broker else '', None),
    (gettext_lazy('Брокер покупателя'), 32, lambda t: t.buyer_broker.name if t.buyer_broker else '', None),
    (gettext_lazy('Число сделок'), 10, lambda t: t.deals_count, INT_FORMAT),
    (gettext_lazy('Объём'), 14, lambda t: t.quantity, QUANTITY_FORMAT),
    (gettext_lazy('Ед. изм.'), 10, lambda t: t.unit, None),
    (gettext_lazy('Цена открытия, тг'), 16, lambda t: t.price_open, PRICE_FORMAT),
    (gettext_lazy('Цена закрытия, тг'), 16, lambda t: t.price_close, PRICE_FORMAT),
    (gettext_lazy('Мин. цена, тг'), 16, lambda t: t.price_min, PRICE_FORMAT),
    (gettext_lazy('Макс. цена, тг'), 16, lambda t: t.price_max, PRICE_FORMAT),
    (gettext_lazy('Средневзв. цена, тг'), 18, lambda t: t.price_avg, PRICE_FORMAT),
    (gettext_lazy('Цена с НДС, тг'), 16, lambda t: t.price_with_vat, PRICE_FORMAT),
    (gettext_lazy('Цена без НДС, тг'), 16, lambda t: t.price_without_vat, PRICE_FORMAT),
    (gettext_lazy('Сумма сделки, тг'), 20, lambda t: t.amount, PRICE_FORMAT),
    (gettext_lazy('Базис поставки'), 22, lambda t: t.delivery_basis, None),
    (gettext_lazy('Место поставки'), 30, lambda t: t.delivery_place, None),
    (gettext_lazy('Срок поставки'), 22, lambda t: t.delivery_term, None),
    (gettext_lazy('Источник'), 40, lambda t: t.source_url, None),
]

# Названия листов (не длиннее 31 символа — ограничение Excel)
SHEET_TRADES = gettext_lazy('Итоги торгов')
SHEET_PARAMS = gettext_lazy('Параметры')


def column_titles() -> list:
    """Заголовки колонок на языке интерфейса."""
    return [str(title) for title, *_rest in COLUMNS]


def _date_format() -> str:
    return DATE_FORMATS.get(i18n.current_language(), DATE_FORMAT)


_FORMAT_STYLES = {}


def _format_style(sheet, number_format):
    """Готовый «стиль ячейки» для числового формата (создаётся один раз на формат).

    Присвоение готового стиля намного быстрее, чем выставление number_format
    в каждой из миллиона ячеек.
    """
    key = (id(sheet.parent), number_format)
    if key not in _FORMAT_STYLES:
        sample = WriteOnlyCell(sheet)
        sample.number_format = number_format
        _FORMAT_STYLES[key] = sample._style
    return _FORMAT_STYLES[key]


def _cell(sheet, value, number_format=None):
    """Значение для записи в строку: простые значения — как есть, с форматом — ячейка со стилем.

    Decimal переводится в float, пустое — в None. Строки, начинающиеся с «=»,
    явно помечаются как текст, чтобы Excel не воспринял данные биржи как формулу.
    """
    if value == '' or value is None:
        return None
    if hasattr(value, 'as_tuple'):  # Decimal
        value = float(value)
    if isinstance(value, str):
        if value[:1] == '=':
            cell = WriteOnlyCell(sheet, value=value)
            cell.data_type = 's'
            return cell
        return value
    if number_format:
        cell = WriteOnlyCell(sheet, value=value)
        cell._style = copy(_format_style(sheet, number_format))
        return cell
    return value


def describe_filters(filters, ordering, exported, truncated, limit):
    """Строки для листа «Параметры»: какие фильтры были применены (на языке интерфейса)."""
    exchange = i18n.exchange_label(filters.exchange) if filters.exchange else _('Обе биржи')

    def day(value):
        return formats.date_format(value, 'SHORT_DATE_FORMAT') if value else ''

    rows = [
        (_('Дата выгрузки'), formats.date_format(timezone.localtime(), 'SHORT_DATETIME_FORMAT')),
        (_('Биржа'), exchange),
        (_('Период с'), day(filters.date_from)),
        (_('Период по'), day(filters.date_to)),
        (_('Товар (ключевое слово)'), filters.q),
        (_('Продавец / покупатель'), filters.party),
        (_('Брокер'), filters.broker),
        (_('Сумма от, тг'), float(filters.amount_min) if filters.amount_min is not None else ''),
        (_('Сумма до, тг'), float(filters.amount_max) if filters.amount_max is not None else ''),
        (_('Место поставки'), filters.delivery),
        (_('Сортировка'), ordering),
        (_('Строк в файле'), exported),
    ]
    if truncated:
        rows.append((
            _('Внимание'),
            _('Выгрузка ограничена %(limit)s строками — уточните фильтры') % {'limit': f'{limit:,}'.replace(',', ' ')},
        ))
    return rows


def build_xlsx(qs, filters, ordering) -> tuple:
    """Строит xlsx по queryset (с select_related). Возвращает (байты, число_строк, усечено)."""
    book = Workbook(write_only=True)
    sheet = book.create_sheet(str(SHEET_TRADES))
    for number, (_title, width, _getter, _fmt) in enumerate(COLUMNS, start=1):
        sheet.column_dimensions[get_column_letter(number)].width = width
    sheet.freeze_panes = 'A2'

    font = Font(bold=True, color='FFFFFF')
    fill = PatternFill('solid', fgColor='1F4E79')
    alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    header = []
    date_format = _date_format()
    for title in column_titles():
        cell = WriteOnlyCell(sheet, value=title)
        cell.font = font
        cell.fill = fill
        cell.alignment = alignment
        header.append(cell)
    sheet.append(header)

    exported = 0
    truncated = False
    # Берём на одну строку больше лимита, чтобы понять, что выборка обрезана
    limit = export_max_rows()
    for t in qs[: limit + 1].iterator(chunk_size=2000):
        if exported >= limit:
            truncated = True
            break
        sheet.append([
            _cell(sheet, getter(t), date_format if fmt == DATE_FORMAT else fmt)
            for _title, _width, getter, fmt in COLUMNS
        ])
        exported += 1
    sheet.auto_filter.ref = f'A1:{get_column_letter(len(COLUMNS))}{max(exported, 1) + 1}'

    params_sheet = book.create_sheet(str(SHEET_PARAMS))
    params_sheet.column_dimensions['A'].width = 28
    params_sheet.column_dimensions['B'].width = 60
    for row in describe_filters(filters, ordering, exported, truncated, limit):
        params_sheet.append([_cell(params_sheet, value) for value in row])

    buffer = BytesIO()
    book.save(buffer)
    return buffer.getvalue(), exported, truncated
