"""Языки интерфейса: русский (по умолчанию), казахский (kk), английский (en).

Проверяется: каждая страница отдаётся на трёх языках (cookie django_language и Accept-Language),
<html lang>, переключатель языка и set_language с возвратом на ту же страницу, переводы
ошибок API, заголовков Excel и подписей статусов, словарь для app.js, полнота файлов .po/.mo.
Данные создаются в тестах; сеть не используется.
"""
import json
import re
import sys
from datetime import date
from io import BytesIO
from pathlib import Path

import polib
import pytest
from django.conf import settings
from django.template.loader import render_to_string
from django.test import Client
from django.utils import timezone, translation
from openpyxl import load_workbook

from exchanges.models import Exchange, News, ParseRun, Product, Trade
from portal import i18n
from portal.js_messages import JS_MESSAGES

pytestmark = pytest.mark.django_db

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / 'scripts'))
import compile_messages  # noqa: E402

LANGUAGES = ('ru', 'kk', 'en')
PAGES = ['/', '/analytics/', '/news/', '/status/',
         '/legal/terms/', '/legal/privacy/', '/legal/cookies/', '/legal/sources/']

# Заголовок <title> каждой страницы на трёх языках
TITLES = {
    '/': ('Итоги торгов — ETS и Каспий', 'Сауда-саттық нәтижелері — ETS және Каспий', 'Trading results — ETS and Caspy'),
    '/analytics/': ('Аналитика — ETS и Каспий', 'Талдау — ETS және Каспий', 'Analytics — ETS and Caspy'),
    '/news/': ('Новости бирж — ETS и Каспий', 'Биржа жаңалықтары — ETS және Каспий', 'Exchange news — ETS and Caspy'),
    '/status/': ('Состояние данных — ETS и Каспий', 'Деректер жағдайы — ETS және Каспий', 'Data status — ETS and Caspy'),
    '/legal/terms/': ('Пользовательское соглашение — ETS и Каспий', 'Пайдаланушы келісімі — ETS және Каспий',
                      'Terms of use — ETS and Caspy'),
    '/legal/privacy/': ('Политика конфиденциальности — ETS и Каспий', 'Құпиялылық саясаты — ETS және Каспий',
                        'Privacy policy — ETS and Caspy'),
    '/legal/cookies/': ('Политика использования cookie — ETS и Каспий', 'Cookie пайдалану саясаты — ETS және Каспий',
                        'Cookie policy — ETS and Caspy'),
    '/legal/sources/': ('Источники данных и лицензии — ETS и Каспий', 'Деректер көздері және лицензиялар — ETS және Каспий',
                        'Data sources and licences — ETS and Caspy'),
}

# Пункты главного меню на трёх языках
MENU = {
    'ru': ['Итоги торгов', 'Аналитика', 'Новости', 'Состояние данных'],
    'kk': ['Сауда-саттық нәтижелері', 'Талдау', 'Жаңалықтар', 'Деректер жағдайы'],
    'en': ['Trading results', 'Analytics', 'News', 'Data status'],
}

TRANSLATION_NOTE = {
    'kk': 'Аударма ыңғайлылық үшін берілген; сәйкессіздік болған жағдайда орыс тіліндегі нұсқасы қолданылады.',
    'en': 'This translation is provided for convenience; in case of any discrepancy, the Russian version prevails.',
}

CYRILLIC = re.compile(r'[А-Яа-яЁёӘәҒғҚқҢңӨөҰұҮүҺһІі]')


def client_for(lang):
    client = Client()
    client.cookies[settings.LANGUAGE_COOKIE_NAME] = lang
    return client


def get_html(lang, address, **extra):
    response = client_for(lang).get(address, **extra)
    assert response.status_code in (200, 404), (lang, address, response.status_code)
    return response.content.decode('utf-8')


def title_of(html):
    match = re.search(r'<title>(.*?)</title>', html, re.S)
    return match.group(1).strip() if match else ''


def i18n_data(html):
    match = re.search(r'<script id="portal-i18n" type="application/json">(.*?)</script>', html, re.S)
    assert match, 'нет словаря portal-i18n'
    return json.loads(match.group(1))


def without_scripts(html):
    return re.sub(r'<script\b.*?</script>', '', html, flags=re.S)


@pytest.fixture
def data():
    ets = Exchange.objects.create(code='ets', name='ETS')
    caspy = Exchange.objects.create(code='caspy', name='Каспий')
    wheat = Product.objects.create(name='Пшеница 3 класса', code='W3')
    Trade.objects.create(exchange=caspy, product=wheat, trade_date=date(2026, 10, 1), price_avg=100,
                         quantity=10, amount=1000, source_key='i18n-1')
    Trade.objects.create(exchange=ets, product=wheat, trade_date=date(2026, 10, 2), price_avg=110,
                         amount=2200, source_key='i18n-2')
    News.objects.create(exchange=caspy, title='Каспий запускает секцию ГСМ', url='https://ccx.kz/news/1',
                        published_at=date(2026, 10, 3))
    ParseRun.objects.create(exchange=ets, status=ParseRun.Status.SUCCESS, trigger='schedule',
                            finished_at=timezone.now())
    return {'ets': ets, 'caspy': caspy, 'wheat': wheat}


# ------------------------------------------------------------ страницы и язык

@pytest.mark.parametrize('lang', LANGUAGES)
@pytest.mark.parametrize('address', PAGES)
def test_every_page_in_three_languages_by_cookie(lang, address):
    html = get_html(lang, address)
    assert f'<html lang="{lang}">' in html
    assert title_of(html) == TITLES[address][LANGUAGES.index(lang)]
    nav = html[html.index('class="main-nav"'):html.index('</nav>')]
    for label in MENU[lang]:
        assert f'>{label}</a>' in nav, (lang, label)
    assert i18n_data(html)['lang'] == lang


@pytest.mark.parametrize('header, lang', [
    ('kk-KZ,kk;q=0.9,ru;q=0.8', 'kk'),
    ('en-US,en;q=0.9', 'en'),
    ('ru-RU,ru;q=0.9', 'ru'),
    ('de-DE,de;q=0.9', 'ru'),   # неподдерживаемый язык — русский
    ('', 'ru'),
])
def test_first_visit_language_from_accept_language(header, lang):
    response = Client().get('/', HTTP_ACCEPT_LANGUAGE=header)
    html = response.content.decode('utf-8')
    assert f'<html lang="{lang}">' in html
    assert title_of(html) == TITLES['/'][LANGUAGES.index(lang)]
    assert 'Accept-Language' in response.get('Vary', '')


def test_cookie_wins_over_accept_language():
    html = get_html('en', '/', HTTP_ACCEPT_LANGUAGE='kk-KZ,kk;q=0.9')
    assert '<html lang="en">' in html


def test_unknown_cookie_value_falls_back_to_russian():
    assert '<html lang="ru">' in get_html('fr', '/')


@pytest.mark.parametrize('lang', LANGUAGES)
def test_404_page_in_three_languages(lang):
    response = client_for(lang).get('/no-such-page/')
    assert response.status_code == 404
    html = response.content.decode('utf-8')
    expected = {'ru': 'Страница не найдена', 'kk': 'Бет табылмады', 'en': 'Page not found'}[lang]
    assert f'<html lang="{lang}">' in html and expected in html
    assert 'class="lang-switch"' in html


@pytest.mark.parametrize('lang, text', [
    ('ru', 'Внутренняя ошибка сервера'), ('kk', 'Сервердің ішкі қатесі'), ('en', 'Internal server error'),
])
def test_500_page_in_three_languages_without_request(lang, text):
    with translation.override(lang):
        html = render_to_string('500.html')
    assert f'<html lang="{lang}">' in html and text in html
    assert i18n_data(html)['lang'] == lang
    assert 'id="theme-toggle"' in html and 'portal/app' in html


def test_english_interface_has_no_russian_text(data):
    """В английском интерфейсе по-русски — только данные бирж и названия языков в переключателе."""
    allowed = [
        'Русский', 'Қазақша',                            # названия языков на самих языках
        'Пшеница 3 класса', 'Каспий запускает секцию ГСМ',  # данные из БД
        'пшеница', 'торги углём',                       # примеры поиска на языке данных
        '«Итоги торгов»', 'Итоги торгов',               # название раздела на сайтах бирж
        '₸',
    ]
    for address in PAGES + ['/no-such-page/']:
        html = without_scripts(get_html('en', address))
        for text in allowed:
            html = html.replace(text, '')
        found = CYRILLIC.search(html)
        assert not found, (address, html[max(0, found.start() - 80):found.start() + 80])


@pytest.mark.parametrize('lang', ('kk', 'en'))
def test_russian_interface_strings_absent_in_other_languages(lang, data):
    for address, strings in {
        '/': ['Скачать Excel', 'Обе биржи', 'Последние новости', 'Строк на странице', 'Сбросить'],
        '/analytics/': ['Оборот по месяцам', 'Сравнение товаров', 'Общая сумма сделок'],
        '/news/': ['Поиск по заголовку', 'Все биржи', 'Найти'],
        '/status/': ['Ручной запуск сбора', 'Последние запуски парсеров'],
    }.items():
        html = without_scripts(get_html(lang, address))
        for text in strings:
            assert text not in html, (lang, address, text)


@pytest.mark.parametrize('lang', ('kk', 'en'))
def test_source_language_note_on_data_pages(lang, data):
    for address in ('/', '/news/', '/analytics/'):
        assert 'class="source-lang-note"' in get_html(lang, address), address
    assert 'class="source-lang-note"' not in get_html('ru', '/')
    # данные (названия товаров, новости) не переводятся
    assert 'Каспий запускает секцию ГСМ' in get_html(lang, '/news/')


@pytest.mark.parametrize('lang', ('kk', 'en'))
@pytest.mark.parametrize('address', [p for p in PAGES if p.startswith('/legal/')])
def test_legal_translations_note_russian_version_prevails(lang, address):
    html = get_html(lang, address)
    assert TRANSLATION_NOTE[lang] in html
    note = html[html.index('id="legal-translation-note"'):]
    note = note[:note.index('</div>')]
    assert 'name="language" value="ru"' in note and 'csrfmiddlewaretoken' in note
    # текст документа — переведённый: русских заголовков разделов нет
    assert '1. Общие положения' not in html and 'Дата редакции:' not in html


@pytest.mark.parametrize('address', [p for p in PAGES if p.startswith('/legal/')])
def test_russian_legal_pages_without_translation_note(address):
    assert 'id="legal-translation-note"' not in get_html('ru', address)


def test_legal_pages_content_translated():
    en = get_html('en', '/legal/privacy/')
    assert 'No. 94-V' in en and 'On Personal Data and Their Protection' in en
    kk = get_html('kk', '/legal/privacy/')
    assert '№ 94-V' in kk and 'Дербес деректер және оларды қорғау туралы' in kk
    assert 'Сайт иесі: АТЫ-ЖӨНІ' in get_html('kk', '/legal/terms/')
    assert 'Site owner: FULL NAME' in get_html('en', '/legal/terms/')


@pytest.mark.parametrize('lang, label', [('ru', 'cookie, необходимая (функциональная)'),
                                         ('kk', 'cookie, қажетті (функционалдық)'),
                                         ('en', 'cookie, necessary (functional)')])
def test_cookie_policy_lists_language_cookie(lang, label):
    html = get_html(lang, '/legal/cookies/')
    row = html[html.index('<code>django_language</code>'):]
    row = row[:row.index('</tr>')]
    assert label in row
    assert str(round(settings.LANGUAGE_COOKIE_AGE / 86400)) in row


# ------------------------------------------------------------ переключатель

@pytest.mark.parametrize('address', PAGES + ['/no-such-page/'])
def test_language_switch_in_header(address):
    html = get_html('kk', address + '?exchange=ets&q=%D1%83%D0%B3%D0%BE%D0%BB%D1%8C')
    header = html[html.index('<header class="site-header">'):html.index('</header>')]
    form = header[header.index('<form class="lang-switch"'):]
    form = form[:form.index('</form>')]
    assert 'action="/i18n/setlang/"' in form and 'method="post"' in form
    assert 'name="csrfmiddlewaretoken"' in form
    # возврат на ту же страницу с тем же query string
    assert f'name="next" value="{address}?exchange=ets&amp;q=%D1%83%D0%B3%D0%BE%D0%BB%D1%8C"' in form
    for code, short in (('ru', 'RU'), ('kk', 'KZ'), ('en', 'EN')):
        assert f'value="{code}" lang="{code}"' in form and f'>{short}</span>' in form
    current = re.findall(r'<button[^>]*aria-current="true"[^>]*>', form)
    assert len(current) == 1 and 'value="kk"' in current[0]
    assert 'id="theme-toggle"' in header


def test_set_language_redirects_back_with_query_and_remembers_choice():
    client = Client()
    response = client.post('/i18n/setlang/', {'language': 'en', 'next': '/news/?exchange=caspy&q=coal'})
    assert response.status_code == 302
    assert response['Location'] == '/news/?exchange=caspy&q=coal'
    cookie = response.cookies[settings.LANGUAGE_COOKIE_NAME]
    assert cookie.value == 'en'
    assert int(cookie['max-age']) == settings.LANGUAGE_COOKIE_AGE >= 300 * 86400
    assert cookie['samesite'] == 'Lax'
    assert '<html lang="en">' in client.get('/news/').content.decode('utf-8')
    # снова на русский
    client.post('/i18n/setlang/', {'language': 'ru', 'next': '/'})
    assert '<html lang="ru">' in client.get('/').content.decode('utf-8')


def test_set_language_requires_csrf_token():
    client = Client(enforce_csrf_checks=True)
    html = client.get('/analytics/').content.decode('utf-8')
    assert client.post('/i18n/setlang/', {'language': 'kk', 'next': '/analytics/'}).status_code == 403
    token = re.search(r'name="csrfmiddlewaretoken" value="([^"]+)"', html).group(1)
    response = client.post('/i18n/setlang/', {'language': 'kk', 'next': '/analytics/', 'csrfmiddlewaretoken': token})
    assert response.status_code == 302 and response['Location'] == '/analytics/'


def test_set_language_rejects_foreign_next_and_unknown_language():
    response = Client().post('/i18n/setlang/', {'language': 'en', 'next': 'https://evil.example/'})
    assert response.status_code == 302 and 'evil.example' not in response['Location']
    # язык не из LANGUAGES: Django может сохранить его в cookie, но страница остаётся русской
    client = Client()
    client.post('/i18n/setlang/', {'language': 'de', 'next': '/'})
    assert '<html lang="ru">' in client.get('/').content.decode('utf-8')


def test_language_switch_script_keeps_current_filters():
    script = (ROOT / 'portal' / 'static' / 'portal' / 'app.js').read_text('utf-8')
    switch = script[script.index('function initLanguageSwitch'):]
    switch = switch[:switch.index('\n  }\n')]
    assert "querySelectorAll('form.lang-switch')" in switch
    assert 'window.location.pathname + window.location.search' in switch


# ------------------------------------------------------------ API и выгрузка

@pytest.mark.parametrize('lang, text', [
    ('ru', 'Параметр date_from: ожидается дата в формате ГГГГ-ММ-ДД'),
    ('kk', 'date_from параметрі: ЖЖЖЖ-АА-КК пішіміндегі күн күтіледі'),
    ('en', 'Parameter date_from: a date in the YYYY-MM-DD format is expected'),
])
def test_api_errors_translated(lang, text):
    response = client_for(lang).get('/api/trades/', {'date_from': 'вчера'})
    assert response.status_code == 400
    assert response.json()['error'] == text


@pytest.mark.parametrize('lang, text', [
    ('kk', 'Салыстыру үшін 10 тауардан артық таңдауға болмайды'),
    ('en', 'You can select no more than 10 products for comparison'),
])
def test_analytics_api_errors_translated(lang, text):
    response = client_for(lang).get('/api/analytics/', {'products': ','.join(str(n) for n in range(1, 12))})
    assert response.status_code == 400 and response.json()['error'].startswith(text)


def test_api_errors_follow_accept_language():
    response = Client().get('/api/trades/', {'ordering': 'nope'}, HTTP_ACCEPT_LANGUAGE='en')
    assert response.json()['error'].startswith('Parameter ordering: invalid field.')


@pytest.mark.parametrize('lang, text', [
    ('kk', 'Демо-нұсқа: деректерді жинау өшірілген.'),
    ('en', 'Demo version: data collection is disabled.'),
])
def test_read_only_collect_message_translated(lang, text, settings):
    settings.PORTAL_READ_ONLY = True
    response = client_for(lang).post('/api/collect/', data=json.dumps({'days': 1}), content_type='application/json')
    assert response.status_code == 403 and response.json()['error'].startswith(text)


@pytest.mark.parametrize('lang, status, trigger, caspy', [
    ('ru', 'Успешно', 'По расписанию', 'Каспий'),
    ('kk', 'Сәтті', 'Кесте бойынша', 'Каспий'),
    ('en', 'Success', 'Scheduled', 'Caspy'),
])
def test_status_labels_and_exchange_names_translated(lang, status, trigger, caspy, data):
    body = client_for(lang).get('/api/status/').json()
    exchanges = {e['code']: e for e in body['exchanges']}
    assert exchanges['ets']['last_run']['status_display'] == status
    assert exchanges['ets']['last_run']['trigger_display'] == trigger
    assert exchanges['ets']['last_run']['status'] == 'success'  # код не меняется
    assert exchanges['caspy']['name'] == caspy
    trades = client_for(lang).get('/api/trades/', {'exchange': 'caspy'}).json()['results']
    assert trades[0]['exchange_name'] == caspy
    assert trades[0]['product'] == 'Пшеница 3 класса'  # данные не переводятся


@pytest.mark.parametrize('lang, sheet, params, header, date_format', [
    ('ru', 'Итоги торгов', 'Параметры', ('Биржа', 'Дата торгов', 'Время', 'Товар'), 'DD.MM.YYYY'),
    ('kk', 'Сауда-саттық нәтижелері', 'Параметрлер', ('Биржа', 'Сауда күні', 'Уақыты', 'Тауар'), 'DD.MM.YYYY'),
    ('en', 'Trading results', 'Parameters', ('Exchange', 'Trade date', 'Time', 'Product'), 'DD/MM/YYYY'),
])
def test_excel_export_in_selected_language(lang, sheet, params, header, date_format, data):
    response = client_for(lang).get('/api/trades/export.xlsx', {'exchange': 'caspy'})
    assert response.status_code == 200
    book = load_workbook(BytesIO(response.content))
    assert book.sheetnames == [sheet, params]
    rows = list(book[sheet].iter_rows(values_only=True))
    assert rows[0][:4] == header
    assert rows[1][0] == ('Caspy' if lang == 'en' else 'Каспий')
    assert rows[1][3] == 'Пшеница 3 класса'
    assert book[sheet]['B2'].number_format == date_format
    labels = [row[0] for row in book[params].iter_rows(values_only=True)]
    expected = {'ru': 'Дата выгрузки', 'kk': 'Жүктеп алу күні', 'en': 'Export date'}[lang]
    assert labels[0] == expected


# ------------------------------------------------------------ словарь для app.js

def test_js_catalog_same_keys_for_all_languages():
    catalogs = {}
    for lang in LANGUAGES:
        catalogs[lang] = i18n_data(get_html(lang, '/'))
        assert catalogs[lang]['lang'] == lang
        assert catalogs[lang]['locale'] == i18n.INTL_LOCALES[lang]
    keys = set(catalogs['ru']['messages'])
    assert keys == set(JS_MESSAGES)
    for lang in LANGUAGES:
        messages = catalogs[lang]['messages']
        assert set(messages) == keys, lang
        assert all(value.strip() for value in messages.values()), lang
    # русский словарь — тождественный; kk/en — переведены
    assert all(catalogs['ru']['messages'][k] == k for k in keys)
    for lang, expected in (('kk', 'Ашық тақырып'), ('en', 'Light theme')):
        assert catalogs[lang]['messages']['Светлая тема'] == expected
    assert catalogs['en']['messages']['янв|фев|мар|апр|май|июн|июл|авг|сен|окт|ноя|дек'].split('|')[9] == 'Oct'
    assert catalogs['kk']['messages']['январь|февраль|март|апрель|май|июнь|июль|август|сентябрь|октябрь|ноябрь|декабрь'].split('|')[9] == 'қазан'
    untranslated = [k for k in keys if catalogs['en']['messages'][k] == k and CYRILLIC.search(k)]
    assert not untranslated, untranslated


def test_js_messages_match_app_js():
    """Список строк для перевода совпадает со строками _('...') и формами слов в app.js."""
    script = (ROOT / 'portal' / 'static' / 'portal' / 'app.js').read_text('utf-8')
    used = set(re.findall(r"\b_\('((?:[^'\\]|\\.)*)'", script))
    used |= set(re.findall(r"(?:const \w+ = |plural\([^,]+, )'([^']*\|[^']*)'", script))
    assert used == set(JS_MESSAGES)
    assert len(JS_MESSAGES) == len(set(JS_MESSAGES))


@pytest.mark.parametrize('lang', ('kk', 'en'))
def test_js_placeholders_and_plural_forms_kept(lang):
    with translation.override(lang):
        catalog = i18n.js_catalog()['messages']
    for key, value in catalog.items():
        assert sorted(re.findall(r'\{(\w+)\}', key)) == sorted(re.findall(r'\{(\w+)\}', value)), key
        if '|' in key and key.count('|') == 11:
            assert value.count('|') == 11, key  # 12 месяцев
        elif '|' in key:
            assert 1 <= len(value.split('|')) <= 2, key  # формы слова: kk — одна, en — две


def test_intl_locale_with_fallback_in_script():
    script = (ROOT / 'portal' / 'static' / 'portal' / 'app.js').read_text('utf-8')
    assert "Intl.NumberFormat.supportedLocalesOf" in script and "'ru-RU'" in script
    assert "new Intl.NumberFormat(LOCALE" in script and "new Intl.DateTimeFormat(LOCALE" in script
    assert "new Intl.NumberFormat('ru-RU'" not in script
    assert i18n.INTL_LOCALES == {'ru': 'ru-RU', 'kk': 'kk-KZ', 'en': 'en-GB'}


# ------------------------------------------------------------ файлы переводов

@pytest.mark.parametrize('lang', ('kk', 'en'))
def test_po_files_complete_without_fuzzy(lang):
    po_path = ROOT / 'locale' / lang / 'LC_MESSAGES' / 'django.po'
    assert compile_messages.check_po(po_path) == []
    po = polib.pofile(str(po_path))
    assert po.percent_translated() == 100
    assert not po.fuzzy_entries() and not po.untranslated_entries()
    assert po.metadata['Language'] == lang


@pytest.mark.parametrize('lang', ('kk', 'en'))
def test_mo_files_compiled_and_up_to_date(lang):
    folder = ROOT / 'locale' / lang / 'LC_MESSAGES'
    po = {(e.msgctxt, e.msgid): e.msgstr for e in polib.pofile(str(folder / 'django.po')).translated_entries()}
    mo = {(e.msgctxt, e.msgid): e.msgstr for e in polib.mofile(str(folder / 'django.mo')) if e.msgid}
    assert po == mo, 'django.mo устарел: python scripts/compile_messages.py'


def test_every_template_string_is_in_po():
    """Все строки {% translate %} шаблонов и gettext в Python есть в файлах переводов."""
    sys.path.insert(0, str(ROOT / 'scripts'))
    import make_messages

    catalog = make_messages.collect()
    for lang in ('kk', 'en'):
        po = polib.pofile(str(ROOT / 'locale' / lang / 'LC_MESSAGES' / 'django.po'))
        known = {(e.msgctxt, e.msgid) for e in po if not e.obsolete}
        missing = [key for key in catalog.items if key not in known]
        assert not missing, (lang, missing[:5], 'python scripts/make_messages.py')


def test_check_detects_problems():
    entry = polib.POEntry(msgid='Найдено: {count}', msgstr='Found: {n}')
    assert compile_messages.check_entry(entry)
    entry = polib.POEntry(msgid='Ошибка', msgstr='', flags=['fuzzy'])
    assert len(compile_messages.check_entry(entry)) == 2
    assert compile_messages.check_entry(polib.POEntry(msgid='Да', msgstr='Yes')) == []


def test_language_settings():
    assert settings.LANGUAGE_CODE == 'ru'
    assert [code for code, _ in settings.LANGUAGES] == ['ru', 'kk', 'en']
    middleware = settings.MIDDLEWARE
    locale = middleware.index('django.middleware.locale.LocaleMiddleware')
    assert middleware.index('django.contrib.sessions.middleware.SessionMiddleware') < locale
    assert locale < middleware.index('django.middleware.common.CommonMiddleware')
    assert Path(settings.LOCALE_PATHS[0]) == ROOT / 'locale'
