"""Тесты портала: JSON API, экспорт в Excel, страницы, ручной запуск сбора.

Данные создаются фабриками прямо в тестах (две биржи, товары по-русски,
участники с БИН, брокеры), сетевые обращения не выполняются.
"""
import json
import sys
import threading
import types
from datetime import date, timedelta
from decimal import Decimal
from io import BytesIO

import pytest
from django.db import connection
from django.test import Client
from django.utils import timezone
from openpyxl import load_workbook

from exchanges.models import Broker, Exchange, Participant, ParseRun, Product, Trade
from portal import collect_runner

pytestmark = pytest.mark.django_db

_счётчик = {'n': 0}


# ------------------------------------------------------------------ фабрики

def создать_биржу(код, название):
    return Exchange.objects.get_or_create(code=код, defaults={'name': название})[0]


def создать_товар(название, код='', tnved=''):
    return Product.objects.get_or_create(name=название, code=код, tnved=tnved)[0]


def создать_участника(название, бин=''):
    return Participant.objects.get_or_create(name=название, bin=бин)[0]


def создать_брокера(название):
    return Broker.objects.get_or_create(name=название)[0]


def создать_сделку(биржа, товар, дата, **поля):
    """Создаёт запись Trade с уникальным source_key."""
    _счётчик['n'] += 1
    поля.setdefault('source_key', f'ключ-{_счётчик["n"]}')
    return Trade.objects.create(exchange=биржа, product=товар, trade_date=дата, **поля)


@pytest.fixture(autouse=True)
def чистое_состояние_сбора():
    """Состояние ручного запуска хранится в памяти процесса — сбрасываем между тестами."""
    collect_runner.wait_finished(2)
    collect_runner.reset_state()
    yield
    collect_runner.wait_finished(2)
    collect_runner.reset_state()


@pytest.fixture
def данные():
    """Набор данных: две биржи, товары, участники с БИН, брокеры, 8 сделок."""
    ets = создать_биржу('ets', 'ETS')
    каспий = создать_биржу('caspy', 'Каспий')

    пшеница = создать_товар('Пшеница мягкая 3 класс', 'WH3', '1001')
    уголь = создать_товар('Уголь каменный марки Д', 'CL1', '2701')
    сахар = создать_товар('Сахар-песок белый', 'SG1', '1701')
    бензин = создать_товар('Бензин АИ-92', 'GS92', '2710')

    агро = создать_участника('ТОО «АгроТрейд»', '123456789012')
    зерновая = создать_участника('ТОО Зерновая компания', '987654321098')
    ктж = создать_участника('АО «Қазақстан Темір Жолы»', '111222333444')
    иванов = создать_участника('ИП Иванов', '')

    алматы = создать_брокера('Брокер Алматы')
    центр = создать_брокера('ЦентрКредит Брокер')

    d = date(2026, 10, 1)
    сделки = {
        'пшеница_ets': создать_сделку(
            ets, пшеница, d + timedelta(days=0), buyer=агро, seller=зерновая,
            buyer_broker=алматы, seller_broker=центр, amount=Decimal('1500000.00'),
            price_avg=Decimal('150000.00'), quantity=Decimal('10.000'), unit='т',
            delivery_place='Костанай', deals_count=2),
        'пшеница_каспий': создать_сделку(
            каспий, пшеница, d + timedelta(days=1), buyer=зерновая, seller=агро,
            buyer_broker=центр, amount=Decimal('3200000.50'),
            price_avg=Decimal('160000.00'), quantity=Decimal('20.000'), unit='т',
            delivery_place='Астана', deals_count=1),
        'уголь_ets': создать_сделку(
            ets, уголь, d + timedelta(days=2), buyer=ктж, seller=иванов,
            amount=Decimal('800000.00'), price_avg=Decimal('8000.00'),
            quantity=Decimal('100.000'), unit='т', delivery_place='Караганда'),
        'уголь_каспий': создать_сделку(
            каспий, уголь, d + timedelta(days=3), buyer=агро, seller=ктж,
            seller_broker=алматы, amount=Decimal('9000000.00'),
            price_avg=Decimal('9000.00'), quantity=Decimal('1000.000'), unit='т',
            delivery_place='Павлодар'),
        'сахар_ets': создать_сделку(
            ets, сахар, d + timedelta(days=4), buyer=иванов, seller=зерновая,
            amount=Decimal('250000.00'), price_avg=Decimal('250.00'),
            delivery_place='Костанай'),
        'сахар_каспий': создать_сделку(
            каспий, сахар, d + timedelta(days=5), buyer=зерновая, seller=иванов,
            amount=None, price_avg=None, delivery_place=''),
        'бензин_ets': создать_сделку(
            ets, бензин, d + timedelta(days=6), buyer=агро, seller=ктж,
            buyer_broker=центр, amount=Decimal('5000000.00'), price_avg=Decimal('200.00'),
            delivery_place='Алматы'),
        'бензин_каспий': создать_сделку(
            каспий, бензин, d + timedelta(days=7), buyer=ктж, seller=агро,
            amount=Decimal('100000.00'), price_avg=Decimal('190.00'),
            delivery_place='Алматы'),
    }
    return types.SimpleNamespace(ets=ets, каспий=каспий, сделки=сделки, d0=d)


@pytest.fixture
def клиент():
    return Client()


def получить(клиент, адрес='/api/trades/', **параметры):
    ответ = клиент.get(адрес, параметры)
    return ответ


def идентификаторы(ответ):
    return [r['id'] for r in ответ.json()['results']]


def ид(данные, *ключи):
    return {данные.сделки[k].id for k in ключи}


# ----------------------------------------------------- поиск по кириллице

def test_rucontains_ищет_кириллицу_без_учёта_регистра(данные):
    найдено = Product.objects.filter(name__rucontains='ПШЕНИЦА')
    assert [p.name for p in найдено] == ['Пшеница мягкая 3 класс']
    assert Product.objects.filter(name__rucontains='пшЕНиЦа МЯГ').count() == 1
    assert Product.objects.filter(name__rucontains='аи-92').count() == 1  # латиница+цифры


def test_rucontains_экранирует_спецсимволы(данные):
    assert Product.objects.filter(name__rucontains='%').count() == 0
    assert Product.objects.filter(name__rucontains='_').count() == 0
    assert Product.objects.filter(name__rucontains='\\').count() == 0


@pytest.mark.skipif(connection.vendor != 'sqlite', reason='особенность SQLite')
def test_встроенный_icontains_на_sqlite_не_знает_регистр_кириллицы(данные):
    """Документирует причину отдельного lookup: icontains на SQLite нечувствителен только к ASCII."""
    assert Product.objects.filter(name__icontains='ПШЕНИЦА').count() == 0
    assert Product.objects.filter(name__rucontains='ПШЕНИЦА').count() == 1


# --------------------------------------------------------- /api/trades/

def test_формат_ответа_и_плоские_поля(клиент, данные):
    ответ = получить(клиент)
    assert ответ.status_code == 200
    assert ответ['Content-Type'].startswith('application/json')
    тело = ответ.json()
    assert set(тело) >= {'count', 'page', 'pages', 'page_size', 'results'}
    assert тело['count'] == 8
    assert тело['page'] == 1 and тело['pages'] == 1 and тело['page_size'] == 50
    строка = next(r for r in тело['results'] if r['id'] == данные.сделки['пшеница_ets'].id)
    assert строка['exchange'] == 'ets'
    assert строка['exchange_name'] == 'ETS'
    assert строка['trade_date'] == '2026-10-01'
    assert строка['product'] == 'Пшеница мягкая 3 класс'
    assert строка['product_code'] == 'WH3'
    assert строка['buyer'] == 'ТОО «АгроТрейд»' and строка['buyer_bin'] == '123456789012'
    assert строка['seller'] == 'ТОО Зерновая компания' and строка['seller_bin'] == '987654321098'
    assert строка['buyer_broker'] == 'Брокер Алматы'
    assert строка['seller_broker'] == 'ЦентрКредит Брокер'
    assert строка['amount'] == 1500000.0
    assert строка['price_avg'] == 150000.0
    assert строка['quantity'] == 10.0 and строка['unit'] == 'т'
    assert строка['delivery_place'] == 'Костанай'
    assert тело['amount_sum'] == pytest.approx(19850000.5)


def test_по_умолчанию_свежие_торги_сверху(клиент, данные):
    даты = [r['trade_date'] for r in получить(клиент).json()['results']]
    assert даты == sorted(даты, reverse=True)
    assert даты[0] == '2026-10-08'


def test_пустые_и_отсутствующие_значения_не_ломают_ответ(клиент, данные):
    строка = next(r for r in получить(клиент).json()['results'] if r['id'] == данные.сделки['сахар_каспий'].id)
    assert строка['amount'] is None and строка['price_avg'] is None
    assert строка['delivery_place'] == ''
    assert строка['trade_time'] is None


@pytest.mark.parametrize('значение, ожидаемые', [
    ('ets', {'пшеница_ets', 'уголь_ets', 'сахар_ets', 'бензин_ets'}),
    ('caspy', {'пшеница_каспий', 'уголь_каспий', 'сахар_каспий', 'бензин_каспий'}),
    ('', None),
    ('ETS', {'пшеница_ets', 'уголь_ets', 'сахар_ets', 'бензин_ets'}),
])
def test_фильтр_биржи(клиент, данные, значение, ожидаемые):
    ответ = получить(клиент, exchange=значение)
    assert ответ.status_code == 200
    if ожидаемые is None:
        assert ответ.json()['count'] == 8
    else:
        assert set(идентификаторы(ответ)) == ид(данные, *ожидаемые)


def test_неизвестная_биржа_даёт_400(клиент, данные):
    ответ = получить(клиент, exchange='nyse')
    assert ответ.status_code == 400
    assert 'error' in ответ.json()


def test_фильтр_периода(клиент, данные):
    ответ = получить(клиент, date_from='2026-10-03', date_to='2026-10-05')
    assert set(идентификаторы(ответ)) == ид(данные, 'уголь_ets', 'уголь_каспий', 'сахар_ets')
    только_от = получить(клиент, date_from='2026-10-07')
    assert set(идентификаторы(только_от)) == ид(данные, 'бензин_ets', 'бензин_каспий')
    только_до = получить(клиент, date_to='2026-10-01')
    assert set(идентификаторы(только_до)) == ид(данные, 'пшеница_ets')


@pytest.mark.parametrize('параметр', ['date_from', 'date_to'])
def test_некорректная_дата_даёт_400(клиент, данные, параметр):
    assert получить(клиент, **{параметр: '31.12.2026'}).status_code == 400


@pytest.mark.parametrize('запрос', ['пшеница', 'ПШЕНИЦА', 'Пшеница', 'пШеНиЦа', 'ПШЕН', 'мягкая 3'])
def test_ключевое_слово_без_учёта_регистра_кириллицы(клиент, данные, запрос):
    ответ = получить(клиент, q=запрос)
    assert set(идентификаторы(ответ)) == ид(данные, 'пшеница_ets', 'пшеница_каспий')


def test_ключевое_слово_несколько_слов_и_несуществующее(клиент, данные):
    assert получить(клиент, q='  уголь   каменный ').json()['count'] == 2
    assert получить(клиент, q='уголь пшеница').json()['count'] == 0
    assert получить(клиент, q='алюминий').json()['count'] == 0
    assert получить(клиент, q='%').json()['count'] == 0


def test_ключевое_слово_латиница_и_цифры(клиент, данные):
    assert set(идентификаторы(получить(клиент, q='аи-92'))) == ид(данные, 'бензин_ets', 'бензин_каспий')


def test_фильтр_участника_по_названию_и_регистру(клиент, данные):
    ответ = получить(клиент, party='агротрейд')
    # АгроТрейд участвует в пяти сделках (покупатель или продавец)
    assert set(идентификаторы(ответ)) == ид(
        данные, 'пшеница_ets', 'пшеница_каспий', 'уголь_каспий', 'бензин_ets', 'бензин_каспий')
    assert получить(клиент, party='АГРОТРЕЙД').json()['count'] == 5
    assert получить(клиент, party='иванов').json()['count'] == 3


def test_фильтр_участника_по_казахским_буквам(клиент, данные):
    ответ = получить(клиент, party='ҚАЗАҚСТАН темір')
    assert set(идентификаторы(ответ)) == ид(данные, 'уголь_ets', 'уголь_каспий', 'бензин_ets', 'бензин_каспий')


def test_фильтр_участника_по_бин(клиент, данные):
    ответ = получить(клиент, party='987654321098')
    assert set(идентификаторы(ответ)) == ид(
        данные, 'пшеница_ets', 'пшеница_каспий', 'сахар_ets', 'сахар_каспий')
    assert получить(клиент, party='111222333444').json()['count'] == 4
    assert получить(клиент, party='000000000000').json()['count'] == 0


def test_фильтр_брокера_покупателя_и_продавца(клиент, данные):
    ответ = получить(клиент, broker='алматы')
    # Брокер Алматы: покупатель в пшеница_ets и продавец в уголь_каспий
    assert set(идентификаторы(ответ)) == ид(данные, 'пшеница_ets', 'уголь_каспий')
    ответ = получить(клиент, broker='ЦЕНТРКРЕДИТ')
    assert set(идентификаторы(ответ)) == ид(данные, 'пшеница_ets', 'пшеница_каспий', 'бензин_ets')


def test_фильтр_суммы(клиент, данные):
    assert set(идентификаторы(получить(клиент, amount_min='5000000'))) == ид(данные, 'уголь_каспий', 'бензин_ets')
    assert set(идентификаторы(получить(клиент, amount_max='250000'))) == ид(данные, 'сахар_ets', 'бензин_каспий')
    диапазон = получить(клиент, amount_min='800 000', amount_max='1500000,00')
    assert set(идентификаторы(диапазон)) == ид(данные, 'уголь_ets', 'пшеница_ets')


def test_записи_без_суммы_не_попадают_в_диапазон(клиент, данные):
    ответ = получить(клиент, amount_min='0')
    assert данные.сделки['сахар_каспий'].id not in идентификаторы(ответ)
    assert ответ.json()['count'] == 7


def test_некорректная_сумма_даёт_400(клиент, данные):
    assert получить(клиент, amount_min='много').status_code == 400
    assert получить(клиент, amount_max='NaN').status_code == 400
    assert получить(клиент, amount_max='1e99').status_code == 400


def test_фильтр_места_поставки(клиент, данные):
    assert set(идентификаторы(получить(клиент, delivery='костанай'))) == ид(данные, 'пшеница_ets', 'сахар_ets')
    assert set(идентификаторы(получить(клиент, delivery='АЛМАТЫ'))) == ид(данные, 'бензин_ets', 'бензин_каспий')
    assert получить(клиент, delivery='Мурманск').json()['count'] == 0


def test_комбинация_фильтров(клиент, данные):
    ответ = получить(
        клиент, exchange='caspy', q='уголь', party='агротрейд',
        date_from='2026-10-01', date_to='2026-10-31', amount_min='1000000', delivery='павлодар',
    )
    assert идентификаторы(ответ) == [данные.сделки['уголь_каспий'].id]
    # те же условия, но другая биржа — пусто
    пусто = получить(клиент, exchange='ets', q='уголь', party='агротрейд')
    assert пусто.json()['count'] == 0
    assert пусто.json()['results'] == [] and пусто.json()['pages'] == 1


def test_комбинация_брокера_и_биржи(клиент, данные):
    ответ = получить(клиент, broker='центркредит', exchange='ets')
    assert set(идентификаторы(ответ)) == ид(данные, 'пшеница_ets', 'бензин_ets')


# -------------------------------------------------------------- сортировка

@pytest.mark.parametrize('поле, ключ', [
    ('trade_date', lambda r: r['trade_date']),
    ('amount', lambda r: r['amount']),
    ('price_avg', lambda r: r['price_avg']),
    ('quantity', lambda r: r['quantity']),
    ('product', lambda r: r['product']),
    ('buyer', lambda r: r['buyer']),
    ('seller', lambda r: r['seller']),
    ('exchange', lambda r: r['exchange_name']),
    ('delivery_place', lambda r: r['delivery_place']),
])
def test_сортировка_по_возрастанию_и_убыванию(клиент, данные, поле, ключ):
    for знак, обратно in (('', False), ('-', True)):
        строки = получить(клиент, ordering=знак + поле).json()['results']
        значения = [ключ(r) for r in строки if ключ(r) not in (None, '')]
        assert значения == sorted(значения, reverse=обратно), f'{знак}{поле}'


def test_пустые_значения_всегда_в_конце(клиент, данные):
    for порядок in ('amount', '-amount'):
        строки = получить(клиент, ordering=порядок).json()['results']
        assert строки[-1]['id'] == данные.сделки['сахар_каспий'].id
        assert all(r['amount'] is not None for r in строки[:-1])


def test_сортировка_по_умолчанию_и_значения_суммы(клиент, данные):
    assert [r['amount'] for r in получить(клиент, ordering='-amount').json()['results']][:3] == [
        9000000.0, 5000000.0, 3200000.5]


@pytest.mark.parametrize('плохое', ['password', '-id', 'exchange__code__x', 'buyer__bin', ';drop', '--'])
def test_недопустимая_сортировка_даёт_400(клиент, данные, плохое):
    assert получить(клиент, ordering=плохое).status_code == 400


def test_сортировка_сохраняется_вместе_с_фильтром(клиент, данные):
    ответ = получить(клиент, exchange='caspy', ordering='-amount')
    суммы = [r['amount'] for r in ответ.json()['results'] if r['amount'] is not None]
    assert суммы == [9000000.0, 3200000.5, 100000.0]
    assert ответ.json()['ordering'] == '-amount'


# --------------------------------------------------------------- пагинация

@pytest.fixture
def много_сделок(данные):
    товар = создать_товар('Пшеница мягкая 3 класс', 'WH3', '1001')
    for i in range(37):
        создать_сделку(данные.ets, товар, date(2026, 9, 1) + timedelta(days=i % 20),
                       amount=Decimal(i + 1) * 1000)
    return данные


def test_пагинация_страницы_и_количество(клиент, много_сделок):
    первая = получить(клиент, page_size=10).json()
    assert (первая['count'], первая['pages'], первая['page'], len(первая['results'])) == (45, 5, 1, 10)
    последняя = получить(клиент, page_size=10, page=5).json()
    assert последняя['page'] == 5 and len(последняя['results']) == 5


def test_страницы_не_пересекаются_и_покрывают_всё(клиент, много_сделок):
    собрано = []
    for номер in range(1, 6):
        собрано += идентификаторы(получить(клиент, page_size=10, page=номер, ordering='trade_date'))
    assert len(собрано) == 45
    assert len(set(собрано)) == 45


def test_размер_страницы_ограничен_10_200(клиент, много_сделок):
    assert получить(клиент, page_size=1).json()['page_size'] == 10
    assert получить(клиент, page_size=5000).json()['page_size'] == 200
    assert получить(клиент, page_size=0).json()['page_size'] == 10
    assert получить(клиент).json()['page_size'] == 50


def test_страница_за_пределом_заменяется_последней(клиент, много_сделок):
    ответ = получить(клиент, page_size=10, page=99).json()
    assert ответ['page'] == 5 and len(ответ['results']) == 5
    assert получить(клиент, page=-3).json()['page'] == 1


def test_некорректные_page_и_page_size_дают_400(клиент, данные):
    assert получить(клиент, page='abc').status_code == 400
    assert получить(клиент, page_size='много').status_code == 400


def test_нет_запросов_на_каждую_строку(клиент, много_сделок, django_assert_max_num_queries):
    # подсчёт, сумма и страница — независимо от числа строк (select_related)
    with django_assert_max_num_queries(3):
        ответ = получить(клиент, page_size=50)
    assert len(ответ.json()['results']) == 45


def test_метод_post_запрещён(клиент, данные):
    assert клиент.post('/api/trades/').status_code == 405


# ----------------------------------------------------------- /api/filters/

def test_api_фильтров_биржи_и_места_поставки(клиент, данные):
    тело = клиент.get('/api/filters/').json()
    коды = [e['code'] for e in тело['exchanges']]
    assert коды[:2] == ['ets', 'caspy']
    места = {m['name']: m['count'] for m in тело['delivery_places']}
    assert места == {'Костанай': 2, 'Алматы': 2, 'Астана': 1, 'Караганда': 1, 'Павлодар': 1}
    assert '' not in места
    assert тело['date_min'] == '2026-10-01' and тело['date_max'] == '2026-10-08'
    # самые частые места идут первыми
    assert {m['name'] for m in тело['delivery_places'][:2]} == {'Костанай', 'Алматы'}


def test_api_фильтров_на_пустой_базе(клиент):
    тело = клиент.get('/api/filters/').json()
    assert {e['code'] for e in тело['exchanges']} == {'ets', 'caspy'}
    assert тело['delivery_places'] == [] and тело['date_min'] is None


@pytest.mark.parametrize('поле, текст, ожидаемое', [
    ('product', 'ПШЕН', ['Пшеница мягкая 3 класс']),
    ('party', 'зерновая', ['ТОО Зерновая компания']),
    ('party', 'ТОО', ['ТОО «АгроТрейд»', 'ТОО Зерновая компания']),
    ('broker', 'центр', ['ЦентрКредит Брокер']),
    ('delivery', 'КОСТ', ['Костанай']),
])
def test_подсказки(клиент, данные, поле, текст, ожидаемое):
    тело = клиент.get('/api/filters/', {'field': поле, 'term': текст}).json()
    assert тело['suggestions'] == ожидаемое


def test_подсказки_короткий_запрос_и_плохое_поле(клиент, данные):
    assert клиент.get('/api/filters/', {'field': 'product', 'term': 'п'}).json()['suggestions'] == []
    assert клиент.get('/api/filters/', {'field': 'secret', 'term': 'пш'}).status_code == 400


# ------------------------------------------------------------ экспорт xlsx

def открыть_xlsx(ответ):
    return load_workbook(BytesIO(ответ.content))


def test_экспорт_xlsx_заголовки_и_строки(клиент, данные):
    ответ = клиент.get('/api/trades/export.xlsx')
    assert ответ.status_code == 200
    assert 'spreadsheetml' in ответ['Content-Type']
    assert 'attachment' in ответ['Content-Disposition'] and '.xlsx' in ответ['Content-Disposition']
    книга = открыть_xlsx(ответ)
    лист = книга['Итоги торгов']
    строки = list(лист.iter_rows(values_only=True))
    заголовок = строки[0]
    assert заголовок[:4] == ('Биржа', 'Дата торгов', 'Время', 'Товар')
    assert 'Сумма сделки, тг' in заголовок and 'Покупатель' in заголовок and 'БИН продавца' in заголовок
    assert len(строки) == 1 + 8
    колонка = {название: i for i, название in enumerate(заголовок)}
    # по умолчанию — свежие сверху
    первая = строки[1]
    assert первая[колонка['Дата торгов']].date() == date(2026, 10, 8)
    assert первая[колонка['Товар']] == 'Бензин АИ-92'
    assert первая[колонка['Биржа']] == 'Каспий'
    assert первая[колонка['Сумма сделки, тг']] == 100000


def test_экспорт_xlsx_форматы_чисел_и_дат(клиент, данные):
    лист = лист_экспорта(клиент, ordering='trade_date', exchange='ets', q='пшеница')
    заголовок = [c.value for c in лист[1]]
    строка = лист[2]
    ячейка_даты = строка[заголовок.index('Дата торгов')]
    assert ячейка_даты.number_format == 'DD.MM.YYYY'
    assert ячейка_даты.value.date() == date(2026, 10, 1)
    сумма = строка[заголовок.index('Сумма сделки, тг')]
    assert сумма.value == 1500000 and сумма.number_format == '#,##0.00'
    assert строка[заголовок.index('Объём')].number_format == '#,##0.000'
    assert строка[заголовок.index('БИН покупателя')].value == '123456789012'
    assert лист.freeze_panes == 'A2'


def лист_экспорта(клиент, **параметры):
    return открыть_xlsx(клиент.get('/api/trades/export.xlsx', параметры))['Итоги торгов']


def test_экспорт_использует_те_же_фильтры_и_сортировку(клиент, данные):
    параметры = dict(exchange='caspy', party='агротрейд', amount_min='100000', ordering='-amount')
    json_суммы = [r['amount'] for r in получить(клиент, **параметры).json()['results']]
    лист = лист_экспорта(клиент, **параметры)
    заголовок = [c.value for c in лист[1]]
    суммы = [row[заголовок.index('Сумма сделки, тг')] for row in лист.iter_rows(min_row=2, values_only=True)]
    assert суммы == json_суммы == [9000000, 3200000.5, 100000]


def test_экспорт_не_зависит_от_page(клиент, данные):
    лист = лист_экспорта(клиент, page='3', page_size='10')
    assert лист.max_row == 9


def test_экспорт_пустой_выборки_содержит_только_заголовок(клиент, данные):
    лист = лист_экспорта(клиент, q='несуществующее')
    assert лист.max_row == 1


def test_экспорт_лист_параметров_описывает_фильтры(клиент, данные):
    книга = открыть_xlsx(клиент.get('/api/trades/export.xlsx', {'q': 'уголь', 'exchange': 'ets', 'date_from': '2026-10-01'}))
    параметры = {a: b for a, b in книга['Параметры'].iter_rows(values_only=True)}
    assert параметры['Биржа'] == 'ETS'
    assert параметры['Товар (ключевое слово)'] == 'уголь'
    assert параметры['Период с'] == '01.10.2026'
    assert параметры['Строк в файле'] == 1


def test_экспорт_ограничен_максимумом_строк(клиент, данные, settings):
    settings.PORTAL_EXPORT_MAX_ROWS = 3
    ответ = клиент.get('/api/trades/export.xlsx')
    книга = открыть_xlsx(ответ)
    assert книга['Итоги торгов'].max_row == 1 + 3
    assert ответ['X-Export-Rows'] == '3' and ответ['X-Export-Truncated'] == '1'
    параметры = {a: b for a, b in книга['Параметры'].iter_rows(values_only=True)}
    assert 'Внимание' in параметры


def test_экспорт_не_исполняет_формулы_из_данных(клиент, данные):
    товар = создать_товар('=1+1', 'F1')
    создать_сделку(данные.ets, товар, date(2026, 10, 20), amount=Decimal('1.00'))
    лист = лист_экспорта(клиент, ordering='-trade_date')
    ячейка = лист[2][3]
    assert ячейка.data_type == 's' and ячейка.value == '=1+1'


def test_экспорт_плохие_параметры_дают_400(клиент, данные):
    assert клиент.get('/api/trades/export.xlsx', {'ordering': 'hack'}).status_code == 400
    assert клиент.get('/api/trades/export.xlsx', {'date_from': 'вчера'}).status_code == 400


# ----------------------------------------------------------------- страницы

def test_главная_страница(клиент, данные):
    ответ = клиент.get('/')
    assert ответ.status_code == 200
    html = ответ.content.decode()
    for ожидаемое in ('Итоги торгов', 'Скачать Excel', 'Сбросить', 'portal/app.js', 'portal/style.css'):
        assert ожидаемое in html
    assert 'csrftoken' in ответ.cookies


def test_страница_состояния_и_api_состояния(клиент, данные):
    assert клиент.get('/status/').status_code == 200
    запуск = ParseRun.objects.create(
        exchange=данные.ets, status=ParseRun.Status.SUCCESS, trigger='schedule',
        date_from=date(2026, 10, 1), date_to=date(2026, 10, 8), rows_added=5, rows_skipped=2,
        finished_at=timezone.now())
    ParseRun.objects.create(
        exchange=данные.каспий, status=ParseRun.Status.FAILED, error='Сайт недоступен',
        finished_at=timezone.now())
    тело = клиент.get('/api/status/').json()
    биржи = {e['code']: e for e in тело['exchanges']}
    assert биржи['ets']['trades_count'] == 4
    assert биржи['ets']['last_trade_date'] == '2026-10-07'
    assert биржи['caspy']['trades_count'] == 4
    assert биржи['caspy']['last_trade_date'] == '2026-10-08'
    assert биржи['ets']['last_run']['id'] == запуск.id
    assert биржи['ets']['last_run']['status'] == 'success'
    assert биржи['ets']['last_run']['status_display'] == 'Успешно'
    assert биржи['ets']['last_run']['rows_added'] == 5 and биржи['ets']['last_run']['rows_skipped'] == 2
    assert биржи['ets']['last_success_at'] is not None
    assert биржи['caspy']['last_run']['error'] == 'Сайт недоступен'
    assert биржи['caspy']['last_success_at'] is None
    assert тело['total_trades'] == 8
    assert len(тело['runs']) == 2
    assert тело['collect']['running'] is False


def test_api_состояния_ограничивает_журнал_десятью_запусками(клиент, данные):
    for _ in range(14):
        ParseRun.objects.create(exchange=данные.ets, status=ParseRun.Status.SUCCESS, finished_at=timezone.now())
    assert len(клиент.get('/api/status/').json()['runs']) == 10


def test_api_состояния_на_пустой_базе(клиент):
    тело = клиент.get('/api/status/').json()
    assert [e['code'] for e in тело['exchanges']] == ['ets', 'caspy']
    assert all(e['trades_count'] == 0 and e['last_run'] is None for e in тело['exchanges'])
    assert тело['runs'] == []


# --------------------------------------------------------- ручной запуск

@pytest.fixture
def подмена_сервиса(monkeypatch):
    """Подставляет scrapers.service.collect: вызовы пишутся, завершение управляется событием."""
    управление = types.SimpleNamespace(
        вызовы=[], начат=threading.Event(), отпустить=threading.Event(), ошибка=None)

    def collect(exchange_codes, date_from, date_to, trigger='manual'):
        управление.вызовы.append((list(exchange_codes), date_from, date_to, trigger))
        управление.начат.set()
        управление.отпустить.wait(5)
        if управление.ошибка:
            raise управление.ошибка

    модуль = types.ModuleType('scrapers.service')
    модуль.collect = collect
    monkeypatch.setitem(sys.modules, 'scrapers.service', модуль)
    yield управление
    управление.отпустить.set()
    collect_runner.wait_finished(3)


def отправить_сбор(клиент, **тело):
    return клиент.post('/api/collect/', data=json.dumps(тело), content_type='application/json')


def test_запуск_сбора_в_потоке_с_параметрами(клиент, данные, подмена_сервиса):
    ответ = отправить_сбор(клиент, exchange='ets', days=5)
    assert ответ.status_code == 202
    тело = ответ.json()
    assert тело['started'] is True and тело['exchanges'] == ['ets']
    сегодня = timezone.localdate()
    assert тело['date_to'] == сегодня.isoformat()
    assert тело['date_from'] == (сегодня - timedelta(days=4)).isoformat()
    assert подмена_сервиса.начат.wait(3)
    коды, с, по, триггер = подмена_сервиса.вызовы[0]
    assert коды == ['ets'] and триггер == 'manual'
    assert (с, по) == (сегодня - timedelta(days=4), сегодня)
    # запрос вернулся, пока сбор ещё «идёт», — страница не зависла
    assert клиент.get('/api/status/').json()['collect']['running'] is True


def test_запуск_обеих_бирж_по_умолчанию(клиент, данные, подмена_сервиса):
    assert отправить_сбор(клиент).status_code == 202
    assert подмена_сервиса.начат.wait(3)
    assert подмена_сервиса.вызовы[0][0] == ['ets', 'caspy']


def test_запуск_с_явным_периодом_и_формой(клиент, данные, подмена_сервиса):
    ответ = клиент.post('/api/collect/', {'exchange': 'caspy', 'date_from': '2026-10-01', 'date_to': '2026-10-05'})
    assert ответ.status_code == 202
    assert подмена_сервиса.начат.wait(3)
    assert подмена_сервиса.вызовы[0] == (['caspy'], date(2026, 10, 1), date(2026, 10, 5), 'manual')


def test_повторный_запуск_пока_идёт_предыдущий_отклоняется(клиент, данные, подмена_сервиса):
    assert отправить_сбор(клиент, days=2).status_code == 202
    assert подмена_сервиса.начат.wait(3)
    повтор = отправить_сбор(клиент, days=2)
    assert повтор.status_code == 409
    assert 'уже выполняется' in повтор.json()['error']
    assert len(подмена_сервиса.вызовы) == 1
    # после завершения запуск снова возможен
    подмена_сервиса.отпустить.set()
    collect_runner.wait_finished(3)
    состояние = клиент.get('/api/status/').json()['collect']
    assert состояние['running'] is False and состояние['finished_at'] and состояние['error'] == ''
    подмена_сервиса.начат.clear()
    assert отправить_сбор(клиент, days=2).status_code == 202


def test_ошибка_в_collect_не_роняет_портал_и_видна_в_состоянии(клиент, данные, подмена_сервиса):
    подмена_сервиса.ошибка = RuntimeError('источник недоступен')
    подмена_сервиса.отпустить.set()
    assert отправить_сбор(клиент).status_code == 202
    collect_runner.wait_finished(3)
    состояние = клиент.get('/api/status/').json()['collect']
    assert состояние['running'] is False
    assert 'источник недоступен' in состояние['error']
    # блокировка снята — можно запустить снова
    assert отправить_сбор(клиент).status_code == 202


def test_свежая_запись_running_в_журнале_блокирует_запуск(клиент, данные, подмена_сервиса):
    ParseRun.objects.create(exchange=данные.ets, status=ParseRun.Status.RUNNING)
    ответ = отправить_сбор(клиент)
    assert ответ.status_code == 409
    assert not подмена_сервиса.вызовы
    assert клиент.get('/api/status/').json()['collect']['running_elsewhere'] is True


def test_зависшая_запись_running_не_блокирует_запуск(клиент, данные, подмена_сервиса):
    запуск = ParseRun.objects.create(exchange=данные.ets, status=ParseRun.Status.RUNNING)
    ParseRun.objects.filter(pk=запуск.pk).update(started_at=timezone.now() - timedelta(hours=3))
    assert отправить_сбор(клиент).status_code == 202


def test_сервис_сбора_недоступен_даёт_503(клиент, данные, monkeypatch):
    monkeypatch.setitem(sys.modules, 'scrapers.service', None)  # импорт вызовет ImportError
    ответ = отправить_сбор(клиент)
    assert ответ.status_code == 503
    assert 'error' in ответ.json()
    assert collect_runner.is_running() is False


@pytest.mark.parametrize('тело', [
    {'exchange': 'nyse'},
    {'exchanges': ['ets', 'xxx']},
    {'days': 0},
    {'days': 5000},
    {'days': 'много'},
    {'date_from': '2026-10-10', 'date_to': '2026-10-01'},
    {'date_from': 'вчера'},
])
def test_некорректные_параметры_запуска_дают_400(клиент, данные, подмена_сервиса, тело):
    assert отправить_сбор(клиент, **тело).status_code == 400
    assert not подмена_сервиса.вызовы


def test_запуск_только_через_post(клиент, данные, подмена_сервиса):
    assert клиент.get('/api/collect/').status_code == 405


def test_запуск_требует_csrf_токен(данные, подмена_сервиса):
    клиент = Client(enforce_csrf_checks=True)
    без_токена = клиент.post('/api/collect/', data='{}', content_type='application/json')
    assert без_токена.status_code == 403
    assert not подмена_сервиса.вызовы
    # токен берём из cookie, как это делает страница
    клиент.get('/status/')
    токен = клиент.cookies['csrftoken'].value
    с_токеном = клиент.post(
        '/api/collect/', data='{"days": 1}', content_type='application/json', HTTP_X_CSRFTOKEN=токен)
    assert с_токеном.status_code == 202
